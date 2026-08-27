"""Node 1: deterministic data extraction.

Walks the traceability chain for a given feature id across the System
Requirements, Command List, and Configuration File workbooks and produces a
single consolidated dict. No LLM calls happen here -- this is pure data
wrangling; later graph nodes are expected to consume `state["extracted"]` for
LLM-driven test case generation.

Every lookup failure is recorded in `state["warnings"]` instead of raising, so
the node always returns its best-effort result alongside a visible list of
what it could not resolve (useful since the sheet layouts here were inferred
from screenshots of a filtered example, not the real workbooks).
"""

from __future__ import annotations

from openpyxl import load_workbook

from src.excel.loader import (
    fill_merged_cells,
    find_cell_containing,
    find_column_by_header,
    find_feature_id_column,
    find_header_row,
    find_row_by_exact_value,
    find_rows_by_exact_value,
    filter_rows_by_checkbox,
    filter_rows_by_marker,
    header_index_map,
    iter_data_rows,
)
from src.graph.state import GraphState

_DEFAULT_HEADER_ROW = 1


def _clean(value) -> str:
    return str(value).strip() if value is not None else ""


def _find_sheet(wb, name_contains: str) -> str | None:
    target = name_contains.lower()
    for name in wb.sheetnames:
        if target in name.lower():
            return name
    return None


def _feature_id_variants(feature_id: str) -> set[str]:
    import re

    raw = str(feature_id).strip()
    variants = {raw}
    digits = re.sub(r"\D", "", raw)
    if digits:
        stripped = digits.lstrip("0") or "0"
        variants.update({digits, stripped, f"#{digits}", f"#{stripped}"})
    return variants


def _looks_like_feature_id_header(header_text: str) -> bool:
    text = header_text.strip()
    if not text:
        return False
    if text.startswith("#"):
        text = text[1:]
    return text.isdigit()


def _strip_marker_columns(record: dict) -> dict:
    return {k: v for k, v in record.items() if not _looks_like_feature_id_header(k)}


def _lookup_feature_name(sysreq_wb, feature_id: str, warnings: list[str]) -> tuple[str | None, dict]:
    sheet_name = _find_sheet(sysreq_wb, "index")
    if not sheet_name:
        warnings.append("Could not find an 'Index' sheet in the System Requirements workbook")
        return None, {}
    ws = sysreq_wb[sheet_name]
    fill_merged_cells(ws)

    hit = None
    for variant in _feature_id_variants(feature_id):
        hit = find_cell_containing(ws, variant)
        if hit:
            break
    if not hit:
        warnings.append(f"Feature id '{feature_id}' not found anywhere in the Index sheet")
        return None, {}

    row_number, _col = hit
    columns = header_index_map(ws, _DEFAULT_HEADER_ROW)
    row_record = {name: ws.cell(row=row_number, column=idx).value for name, idx in columns.items()}

    feature_name = None
    for header_name, idx in columns.items():
        if "feature" in header_name and "name" in header_name:
            feature_name = ws.cell(row=row_number, column=idx).value
            break
    if feature_name is None:
        for header_name, idx in columns.items():
            if "name" in header_name:
                feature_name = ws.cell(row=row_number, column=idx).value
                break
    if feature_name is None:
        warnings.append(
            f"Found feature id '{feature_id}' in Index sheet row {row_number} but could not identify "
            "a 'feature name' column - returning the raw row instead"
        )

    return (_clean(feature_name) or None), row_record


def _extract_requirements(sysreq_wb, feature_id: str, warnings: list[str]) -> dict:
    if feature_id not in sysreq_wb.sheetnames:
        warnings.append(f"No requirement sheet named '{feature_id}' found in System Requirements workbook")
        return {"metadata": [], "functional_requirements": []}

    ws = sysreq_wb[feature_id]
    fill_merged_cells(ws)
    try:
        header_row = find_header_row(ws, ["requirement id"])
    except ValueError:
        header_row = _DEFAULT_HEADER_ROW

    metadata = []
    functional_requirements = []
    for _row_number, record in iter_data_rows(ws, header_row):
        category = _clean(record.get("category"))
        if not category:
            continue
        if category.lower() == "functional requirement":
            functional_requirements.append(record)
        else:
            metadata.append(record)

    if not functional_requirements:
        warnings.append(f"No rows with Category == 'Functional Requirement' found in sheet '{feature_id}'")

    return {"metadata": metadata, "functional_requirements": functional_requirements}


def _extract_master_sheet(sysreq_wb, name_contains: str, feature_id: str, warnings: list[str]) -> list[dict]:
    sheet_name = _find_sheet(sysreq_wb, name_contains)
    if not sheet_name:
        warnings.append(f"Could not find a sheet matching '{name_contains}' in System Requirements workbook")
        return []

    ws = sysreq_wb[sheet_name]
    fill_merged_cells(ws)
    col_idx = find_feature_id_column(ws, _DEFAULT_HEADER_ROW, feature_id)
    if col_idx is None:
        warnings.append(f"Could not find a column for feature id '{feature_id}' in sheet '{sheet_name}'")
        return []

    results = [
        _strip_marker_columns(record)
        for _row_number, record in filter_rows_by_marker(ws, _DEFAULT_HEADER_ROW, col_idx, {"o"})
    ]
    if not results:
        warnings.append(f"No rows marked 'O' for feature id '{feature_id}' in sheet '{sheet_name}'")
    return results


def _lookup_numeric_feature_id(command_wb, feature_name: str | None, warnings: list[str]) -> str | None:
    if not feature_name:
        warnings.append("Skipping Feature_ID_Mapping lookup because the feature name was not resolved")
        return None

    sheet_name = _find_sheet(command_wb, "feature_id_mapping") or _find_sheet(command_wb, "feature id mapping")
    if not sheet_name:
        warnings.append("Could not find a Feature_ID_Mapping sheet in the Command List workbook")
        return None

    ws = command_wb[sheet_name]
    fill_merged_cells(ws)
    record = find_row_by_exact_value(ws, _DEFAULT_HEADER_ROW, "Feature Name", feature_name)
    if not record:
        warnings.append(f"Feature name '{feature_name}' not found in Feature_ID_Mapping sheet")
        return None

    numeric_id = record.get("id")
    if numeric_id is None:
        warnings.append(f"Feature_ID_Mapping row for '{feature_name}' has no 'ID' column value")
        return None
    return _clean(numeric_id)


def _extract_feature_signals(
    command_wb, name_contains: str, numeric_feature_id: str | None, warnings: list[str]
) -> tuple[list[dict], set[str]]:
    sheet_name = _find_sheet(command_wb, name_contains)
    if not sheet_name:
        warnings.append(f"Could not find a sheet matching '{name_contains}' in the Command List workbook")
        return [], set()

    if not numeric_feature_id:
        warnings.append(f"Skipping '{sheet_name}' filtering because the numeric feature id was not resolved")
        return [], set()

    ws = command_wb[sheet_name]
    fill_merged_cells(ws)
    col_idx = find_feature_id_column(ws, _DEFAULT_HEADER_ROW, numeric_feature_id)
    if col_idx is None:
        warnings.append(f"Could not find column '{numeric_feature_id}' in sheet '{sheet_name}'")
        return [], set()

    rows = []
    command_names: set[str] = set()
    for _row_number, record in filter_rows_by_checkbox(ws, _DEFAULT_HEADER_ROW, col_idx):
        rows.append(_strip_marker_columns(record))
        cmd_name = _clean(record.get("command name"))
        if cmd_name:
            command_names.add(cmd_name)

    if not rows:
        warnings.append(
            f"No checked rows found in '{sheet_name}' for feature id '{numeric_feature_id}' "
            "(if the real workbook uses unlinked Excel checkbox controls rather than TRUE/FALSE "
            "cell values, this filter will need to be revisited)"
        )
    return rows, command_names


def _extract_command_details(command_wb, config_wb, command_names: set[str], warnings: list[str]) -> list[dict]:
    command_sheet_name = _find_sheet(command_wb, "command list")
    command_ws = None
    if not command_sheet_name:
        warnings.append("Could not find a 'Command List' sheet in the Command List workbook")
    else:
        command_ws = command_wb[command_sheet_name]
        fill_merged_cells(command_ws)

    model_sheet_name = _find_sheet(config_wb, "model_input_mapping") or _find_sheet(config_wb, "model input mapping")
    model_ws = None
    if not model_sheet_name:
        warnings.append("Could not find a 'Model_Input_Mapping' sheet in the Configuration File workbook")
    else:
        model_ws = config_wb[model_sheet_name]
        fill_merged_cells(model_ws)

    commands = []
    for command_name in sorted(command_names):
        entry: dict = {"command_name": command_name, "command_list_details": None, "model_input_mapping": []}

        details = None
        if command_ws is not None:
            details = find_row_by_exact_value(command_ws, _DEFAULT_HEADER_ROW, "Command name", command_name)
            if details is None:
                warnings.append(f"Command '{command_name}' not found in the Command List sheet")
        entry["command_list_details"] = details

        signal_name = _clean(details.get("signal name")) if details else ""
        if model_ws is not None:
            if signal_name:
                matches = find_rows_by_exact_value(model_ws, _DEFAULT_HEADER_ROW, "Signal", signal_name)
                if not matches:
                    warnings.append(
                        f"No Model_Input_Mapping match for signal '{signal_name}' (command '{command_name}')"
                    )
                entry["model_input_mapping"] = matches
            else:
                warnings.append(
                    f"Command '{command_name}' has no Signal Name to match against Model_Input_Mapping"
                )

        commands.append(entry)
    return commands


def extract_data(state: GraphState) -> GraphState:
    warnings: list[str] = []
    feature_id = state["feature_id"]

    sysreq_wb = load_workbook(state["sysreq_path"], data_only=True)
    command_wb = load_workbook(state["command_list_path"], data_only=True)
    config_wb = load_workbook(state["config_path"], data_only=True)

    feature_name, index_row = _lookup_feature_name(sysreq_wb, feature_id, warnings)
    requirements = _extract_requirements(sysreq_wb, feature_id, warnings)
    comm_matrix = _extract_master_sheet(sysreq_wb, "master comm matrix", feature_id, warnings)
    app_parameters = _extract_master_sheet(sysreq_wb, "app parameter", feature_id, warnings)
    io_signals = _extract_master_sheet(sysreq_wb, "input output signals", feature_id, warnings)

    numeric_feature_id = _lookup_numeric_feature_id(command_wb, feature_name, warnings)

    recorder_signals, recorder_commands = _extract_feature_signals(
        command_wb, "recorder_signals_feature", numeric_feature_id, warnings
    )
    sdo_signals, sdo_commands = _extract_feature_signals(
        command_wb, "sdo_signals_feature", numeric_feature_id, warnings
    )

    command_names = recorder_commands | sdo_commands
    commands = _extract_command_details(command_wb, config_wb, command_names, warnings)

    extracted = {
        "feature_id": feature_id,
        "feature_name": feature_name,
        "index_row": index_row,
        "requirements": requirements,
        "comm_matrix": comm_matrix,
        "app_parameters": app_parameters,
        "io_signals": io_signals,
        "numeric_feature_id": numeric_feature_id,
        "recorder_signals": recorder_signals,
        "sdo_signals": sdo_signals,
        "commands": commands,
    }

    return {**state, "extracted": extracted, "warnings": warnings}
