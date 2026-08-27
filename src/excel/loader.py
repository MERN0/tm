"""Generic, header-text-driven helpers for reading traceability data out of the
project's Excel workbooks.

These helpers deliberately avoid hard-coding column positions. Column layout in
the real workbooks was not observed directly (only screenshots of a filtered
example were available), so every lookup is done by matching header text or by
scanning for a target value, which keeps the extraction node tolerant of
reordered/renamed columns in the real files.
"""

from __future__ import annotations

import re
from typing import Iterable

from openpyxl.worksheet.worksheet import Worksheet

MAX_HEADER_SCAN_ROWS = 5

# Cell values that should be treated as "checked" for boolean/checkbox-style columns.
_TRUTHY_CHECKBOX_VALUES = {"true", "1", "x", "yes", "checked", "o"}


def _normalize(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_lower(value) -> str:
    return _normalize(value).lower()


def find_header_row(ws: Worksheet, must_contain: Iterable[str], max_scan_rows: int = MAX_HEADER_SCAN_ROWS) -> int:
    """Return the 1-indexed row number of the header row: the first row (within
    the first `max_scan_rows`) whose cell text contains all of `must_contain`
    (case-insensitive substring match, any column).

    Raises ValueError if no matching row is found.
    """
    wanted = [s.lower() for s in must_contain]
    for row in range(1, max_scan_rows + 1):
        row_text = " ".join(_normalize_lower(c.value) for c in ws[row])
        if all(token in row_text for token in wanted):
            return row
    raise ValueError(f"Could not locate header row containing {list(must_contain)!r} in sheet '{ws.title}'")


def header_index_map(ws: Worksheet, header_row: int) -> dict[str, int]:
    """Map normalized (lowercased, stripped) header text -> 1-indexed column number."""
    mapping: dict[str, int] = {}
    for cell in ws[header_row]:
        text = _normalize_lower(cell.value)
        if text:
            mapping[text] = cell.column
    return mapping


def find_column_by_header(ws: Worksheet, header_row: int, header_text: str) -> int | None:
    """Exact (case-insensitive, stripped) header match -> 1-indexed column number, or None."""
    target = header_text.strip().lower()
    for cell in ws[header_row]:
        if _normalize_lower(cell.value) == target:
            return cell.column
    return None


def _feature_id_variants(feature_id: str) -> set[str]:
    """Build the set of textual forms a feature id might appear as in a header
    cell, e.g. "019" -> {"019", "19", "#19", "#019"}."""
    raw = str(feature_id).strip()
    variants = {raw}
    digits = re.sub(r"\D", "", raw)
    if digits:
        variants.add(digits)
        variants.add(digits.lstrip("0") or "0")
        variants.add(f"#{digits}")
        variants.add(f"#{digits.lstrip('0') or '0'}")
    return {v.lower() for v in variants}


def find_feature_id_column(ws: Worksheet, header_row: int, feature_id: str) -> int | None:
    """Find the 1-indexed column whose header matches the given feature id,
    tolerating "019" / "19" / "#3" style variants. Returns None if not found."""
    variants = _feature_id_variants(feature_id)
    for cell in ws[header_row]:
        text = _normalize_lower(cell.value)
        if text in variants:
            return cell.column
    return None


def iter_data_rows(ws: Worksheet, header_row: int):
    """Yield (row_number, {normalized_header: cell_value}) for every non-empty
    row after the header row."""
    columns = header_index_map(ws, header_row)
    for row_number in range(header_row + 1, ws.max_row + 1):
        row_cells = ws[row_number]
        if all(c.value is None for c in row_cells):
            continue
        record = {name: ws.cell(row=row_number, column=col_idx).value for name, col_idx in columns.items()}
        yield row_number, record


def filter_rows_by_marker(
    ws: Worksheet,
    header_row: int,
    marker_col_idx: int,
    valid_markers: set[str] | None = None,
):
    """Yield (row_number, record_dict) for rows where the cell in
    `marker_col_idx` (case-insensitive, stripped) is in `valid_markers`
    (default: {"o"})."""
    valid = valid_markers or {"o"}
    for row_number, record in iter_data_rows(ws, header_row):
        cell_value = _normalize_lower(ws.cell(row=row_number, column=marker_col_idx).value)
        if cell_value in valid:
            yield row_number, record


def filter_rows_by_checkbox(ws: Worksheet, header_row: int, checkbox_col_idx: int):
    """Yield (row_number, record_dict) for rows where the cell in
    `checkbox_col_idx` looks truthy/checked (TRUE, 1, "x", "checked", ...).

    Note: this only works if the checkbox is linked to the cell's value.
    Unlinked Excel form-control checkboxes are drawing objects openpyxl cannot
    read, and will silently look empty here -- callers should treat an
    all-empty result as a signal to verify against the real workbook.
    """
    for row_number, record in iter_data_rows(ws, header_row):
        raw = ws.cell(row=row_number, column=checkbox_col_idx).value
        if isinstance(raw, bool):
            checked = raw
        else:
            checked = _normalize_lower(raw) in _TRUTHY_CHECKBOX_VALUES
        if checked:
            yield row_number, record


def find_row_by_exact_value(ws: Worksheet, header_row: int, col_name: str, value: str) -> dict | None:
    """Return the first row (as a record dict) where the column named
    `col_name` (case-insensitive header match) equals `value` exactly
    (stripped, case-sensitive comparison on content). Returns None if not found."""
    col_idx = find_column_by_header(ws, header_row, col_name)
    if col_idx is None:
        return None
    target = _normalize(value)
    for row_number, record in iter_data_rows(ws, header_row):
        cell_value = _normalize(ws.cell(row=row_number, column=col_idx).value)
        if cell_value == target:
            return record
    return None


def find_rows_by_exact_value(ws: Worksheet, header_row: int, col_name: str, value: str) -> list[dict]:
    """Like find_row_by_exact_value but returns every matching row."""
    col_idx = find_column_by_header(ws, header_row, col_name)
    if col_idx is None:
        return []
    target = _normalize(value)
    return [
        record
        for row_number, record in iter_data_rows(ws, header_row)
        if _normalize(ws.cell(row=row_number, column=col_idx).value) == target
    ]


def fill_merged_cells(ws: Worksheet) -> None:
    """Unmerge every merged range in the sheet and copy the top-left cell's
    value into every cell the range covered.

    Grouped rows (e.g. a repeated Signal/heading spanning several rows in the
    real workbooks) are commonly represented with merged cells in Excel; to
    openpyxl every cell but the top-left of a merged range reads as None, which
    would otherwise make row-by-row extraction silently lose that value on all
    but the first row of the group. This mutates the in-memory worksheet only
    (the source file on disk is never written to).
    """
    for merged_range in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left_value = ws.cell(row=min_row, column=min_col).value
        ws.unmerge_cells(str(merged_range))
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                ws.cell(row=row, column=col).value = top_left_value


def find_cell_containing(ws: Worksheet, text: str, max_rows: int | None = None, max_cols: int | None = None):
    """Scan the sheet for the first cell whose stripped text exactly equals
    `text` (case-insensitive). Returns (row, col) or None."""
    target = text.strip().lower()
    max_rows = max_rows or ws.max_row
    max_cols = max_cols or ws.max_column
    for row in ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols):
        for cell in row:
            if _normalize_lower(cell.value) == target:
                return cell.row, cell.column
    return None
